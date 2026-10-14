# -*- coding: utf-8 -*-
"""담당자 수정본 양식(데이터_수정본양식.xlsx) → 표준 양식 v2.

수정본의 열 구성·필수(노랑) 표시는 그대로 두고 다음만 고친다.
  - 3_주택연금 '대출잔액'을 입력 칸(노랑)으로: 지급·이자·비용 누계 열이 빠져 수식이 #REF! 였음
  - 1_주택담보대출 LTV 수식을 '배분익스포저 ÷ 평가금액'으로 정정, 주택연금 비율 수식 유지
  - 드롭다운(코드표 목록)·자릿수 검증 다시 걸기 (openpyxl 이 수정본의 확장 검증을 읽지 못함)
  - 데이터는 지우고 예시 3행만 회색 기울임으로 남김, 작성안내 문구를 현재 열 구성에 맞춤

실행: python3 make_template_v2.py ../샘플데이터/데이터_수정본양식.xlsx ../샘플데이터/부동산통합_로우데이터_양식_v2.xlsx
"""
import sys
import warnings
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings("ignore")
src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
REQ = PatternFill("solid", fgColor="FFE699")
EX = Font(italic=True, color="808080")
MAX = 3003


def codes(ws):
    return [str(c.value).split("\n")[-1].strip() if c.value else None for c in ws[3]]


def col(ws, code):
    return get_column_letter(codes(ws).index(code) + 1)


for name in ("1_주택담보대출", "2_투자부동산", "3_주택연금"):
    ws = wb[name]
    cs = codes(ws)
    for r in range(4, ws.max_row + 1):                      # 데이터 비우기 (수식 열은 유지), 예시 3행은 회색
        for j, code in enumerate(cs, start=1):
            c = ws.cell(row=r, column=j)
            if r <= 6:
                c.font = EX
            elif not str(c.value or "").startswith("="):
                c.value = None
    ws.data_validations.dataValidation = []

ws = wb["1_주택담보대출"]
E, A, L = col(ws, "EXPO_AMT"), col(ws, "APPR_AMT"), col(ws, "LTV")
for r in range(4, MAX + 1):
    ws[f"{L}{r}"] = f'=IF(OR({A}{r}="",{A}{r}=0,{E}{r}=""),"",{E}{r}/{A}{r})'

ws = wb["3_주택연금"]
E = col(ws, "EXPO_AMT")
h = ws[f"{E}3"]
h.value, h.fill = "대출잔액(원)\nEXPO_AMT", REQ
for r in range(4, MAX + 1):
    if str(ws[f"{E}{r}"].value or "").startswith("="):
        ws[f"{E}{r}"] = None
ex = {4: 186000000, 5: 233000000, 6: 290000000}
for r, v in ex.items():
    if ws.cell(row=r, column=1).value:
        ws[f"{E}{r}"] = v
        ws[f"{E}{r}"].font = EX

# 드롭다운·자릿수 검증
code_ws = wb["코드표"]
heads = [c.value for c in code_ws[1]]
LISTS = {"RE_TYPE": "부동산유형", "CPTY_TYPE": "차주구분", "ASSET_QLTY": "건전성분류", "APPR_SRC": "평가구분", "PROD": "상품구분",
         "ACCT": "계정구분", "HOLD_TYPE": "보유형태", "PAY_TYPE": "지급방식"}
for name in ("1_주택담보대출", "2_투자부동산", "3_주택연금"):
    ws = wb[name]
    cs = codes(ws)
    for code, key in LISTS.items():
        if code not in cs or key not in heads:
            continue
        k = get_column_letter(heads.index(key) + 1)
        n = sum(1 for c in code_ws[k][1:] if c.value not in (None, ""))
        dv = DataValidation(type="list", formula1=f"=코드표!${k}$2:${k}${n + 1}", allow_blank=True, error="코드표에 있는 값만 입력하세요.", errorTitle="코드 오류")
        dv.add(f"{col(ws, code)}4:{col(ws, code)}{MAX}")
        ws.add_data_validation(dv)
    dv = DataValidation(type="textLength", operator="equal", formula1="6", allow_blank=True, error="YYYYMM 6자리로 입력하세요.", errorTitle="자릿수 오류")
    dv.add(f"{col(ws, 'BASE_YM')}4:{col(ws, 'BASE_YM')}{MAX}")
    ws.add_data_validation(dv)

# 작성안내 갱신
g = wb["작성안내"]
for row in g.iter_rows():
    for c in row:
        c.value = None
lines = [
    ("부동산 통합관리 로우데이터 작성안내 (v2)", True),
    ("", False),
    ("1. 시트 구성", True),
    ("• 1_주택담보대출 / 2_투자부동산 / 3_주택연금: 원천(SAP·부서자료)별 로우데이터. 1행 = 익스포저 × 부동산 1건.", False),
    ("• 4_가격지표: 시군구별 월간 매매가격지수(선택 — 시스템의 '가격지표 월별 업데이트'로 따로 올려도 됨).", False),
    ("• 코드표: 드롭다운 목록 원본.", False),
    ("", False),
    ("2. 헤더 색상", True),
    ("• 노랑 = 필수 / 파랑 = 해당 시 입력 / 초록·주황 = 자동계산(입력 금지). 시스템은 노랑 칸이 빈 행을 '필수항목 누락'으로 표시합니다.", False),
    ("", False),
    ("3. 공통 규칙", True),
    ("• 기준년월 BASE_YM: YYYYMM 6자리 텍스트 (예: 202609).", False),
    ("• 관리번호: 주담대·투자부동산은 관리번호(EXPO_ID), 주택연금은 계약번호(SRC_NO)를 씁니다. 기준월이 바뀌어도 같은 건은 같은 번호.", False),
    ("• 부동산순번 RE_SEQ: 같은 관리번호 안에서 1부터. 공동담보는 물건마다 1행.", False),
    ("• 도로명주소: 시도 시군구 도로명 건물번호까지. 공동주택은 괄호 참고항목 (법정동, 단지명)을 붙이면 주소DB 없이도 동을 정할 수 있습니다. 동·호수는 넣지 않아도 됩니다.", False),
    ("• 법정동코드·좌표는 넣지 않습니다. 시스템이 도로명주소로 붙입니다.", False),
    ("• 금액은 원 단위 정수, 날짜는 날짜형.", False),
    ("", False),
    ("4. 원천별 익스포저 산식 (시스템이 다시 계산)", True),
    ("• 주택담보대출: 배분익스포저 = 대출잔액 × 배분비율. 관리번호별 배분비율 합 = 100%. LTV = 배분익스포저 ÷ 평가금액.", False),
    ("• 투자부동산: 익스포저 = 공정가치 × 지분율.", False),
    ("• 주택연금: 익스포저 = 대출잔액(직접 입력). 잔액 ÷ 주택가격이 100%에 가까울수록 담보부족 위험.", False),
    ("", False),
    ("5. 개인정보", True),
    ("• 차주 실명·주민번호·연락처는 넣지 않습니다. 내부 고객번호만 사용합니다.", False),
    ("", False),
    ("※ 회색 기울임 예시 행은 지우고 사용하세요(수식 열은 남겨둡니다).", False),
]
for i, (t, b) in enumerate(lines, start=1):
    g.cell(row=i, column=1, value=t).font = Font(bold=b, size=13 if i == 1 else 10)
wb.save(out)
print("saved", out)
